import json
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


UNIDADES = [
    'zero', 'um', 'dois', 'três', 'quatro',
    'cinco', 'seis', 'sete', 'oito', 'nove'
]

DEZ_A_DEZENOVE = [
    'dez', 'onze', 'doze', 'treze', 'quatorze',
    'quinze', 'dezesseis', 'dezessete', 'dezoito', 'dezenove'
]

DEZENAS = [
    '', '', 'vinte', 'trinta', 'quarenta',
    'cinquenta', 'sessenta', 'setenta', 'oitenta', 'noventa'
]

CENTENAS = [
    '', 'cento', 'duzentos', 'trezentos', 'quatrocentos',
    'quinhentos', 'seiscentos', 'setecentos', 'oitocentos', 'novecentos'
]


def extenso_ate_999(numero):
    numero = int(numero)

    if numero < 10:
        return UNIDADES[numero]

    if numero < 20:
        return DEZ_A_DEZENOVE[numero - 10]

    if numero < 100:
        dezena, unidade = divmod(numero, 10)
        texto = DEZENAS[dezena]

        if unidade:
            texto += ' e ' + UNIDADES[unidade]

        return texto

    if numero == 100:
        return 'cem'

    centena, resto = divmod(numero, 100)
    texto = CENTENAS[centena]

    if resto:
        texto += ' e ' + extenso_ate_999(resto)

    return texto


def extenso_inteiro(numero):
    numero = int(numero)

    if numero == 0:
        return 'zero'

    if numero < 1000:
        return extenso_ate_999(numero)

    grupos = []

    milhoes, resto = divmod(numero, 1_000_000)

    if milhoes:
        if milhoes == 1:
            grupos.append('um milhão')
        else:
            grupos.append(extenso_ate_999(milhoes) + ' milhões')

    milhares, unidades = divmod(resto, 1000)

    if milhares:
        if milhares == 1:
            grupos.append('mil')
        else:
            grupos.append(extenso_ate_999(milhares) + ' mil')

    if unidades:
        texto_unidades = extenso_ate_999(unidades)

        if grupos and (unidades < 100 or unidades % 100 == 0):
            return ' '.join(grupos) + ' e ' + texto_unidades

        grupos.append(texto_unidades)

    return ' '.join(grupos)


def valor_por_extenso(valor):
    valor = Decimal(str(valor)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    reais = int(valor)
    centavos = int((valor - Decimal(reais)) * 100)

    partes = []

    if reais:
        texto_reais = extenso_inteiro(reais)
        texto_reais += ' real' if reais == 1 else ' reais'
        partes.append(texto_reais)

    if centavos:
        texto_centavos = extenso_inteiro(centavos)
        texto_centavos += ' centavo' if centavos == 1 else ' centavos'
        partes.append(texto_centavos)

    if not partes:
        return 'zero reais'

    return ' e '.join(partes)


def remover_paragrafo(paragrafo):
    elemento = paragrafo._element
    pai = elemento.getparent()

    if pai is not None:
        pai.remove(elemento)


def substituir_em_paragrafo(paragrafo, referencias):
    for codigo, valor in referencias.items():
        valor = str(valor)

        if codigo not in paragrafo.text:
            continue

        if codigo == '{{CLAUSULAS_EXTRAS}}':
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragrafo.paragraph_format.left_indent = Pt(0)
            paragrafo.paragraph_format.first_line_indent = Pt(0)

        substituido = False

        for run in paragrafo.runs:
            if codigo in run.text:
                run.text = run.text.replace(codigo, valor)
                substituido = True

        if substituido:
            continue

        # Fallback para placeholders eventualmente divididos entre vários runs.
        texto = paragrafo.text.replace(codigo, valor)

        if paragrafo.runs:
            paragrafo.runs[0].text = texto

            for run in paragrafo.runs[1:]:
                run.text = ''
        else:
            paragrafo.add_run(texto)


def processar_paragrafos(paragrafos, referencias, remover_se_vazio):
    for paragrafo in list(paragrafos):
        texto = paragrafo.text.strip()

        if texto in remover_se_vazio and not str(referencias.get(texto, '')).strip():
            remover_paragrafo(paragrafo)
            continue

        substituir_em_paragrafo(paragrafo, referencias)


def processar_tabela(tabela, referencias, remover_se_vazio):
    for linha in tabela.rows:
        for celula in linha.cells:
            processar_paragrafos(celula.paragraphs, referencias, remover_se_vazio)

            for tabela_interna in celula.tables:
                processar_tabela(tabela_interna, referencias, remover_se_vazio)


def processar_container(container, referencias, remover_se_vazio):
    processar_paragrafos(container.paragraphs, referencias, remover_se_vazio)

    for tabela in container.tables:
        processar_tabela(tabela, referencias, remover_se_vazio)


def coletar_texto_documento(documento):
    partes = []

    for paragrafo in documento.paragraphs:
        partes.append(paragrafo.text)

    def coletar_tabela(tabela):
        for linha in tabela.rows:
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    partes.append(paragrafo.text)

                for tabela_interna in celula.tables:
                    coletar_tabela(tabela_interna)

    for tabela in documento.tables:
        coletar_tabela(tabela)

    for secao in documento.sections:
        for container in [
            secao.header,
            secao.first_page_header,
            secao.even_page_header,
            secao.footer,
            secao.first_page_footer,
            secao.even_page_footer,
        ]:
            for paragrafo in container.paragraphs:
                partes.append(paragrafo.text)

            for tabela in container.tables:
                coletar_tabela(tabela)

    return '\n'.join(partes)


def gerar_contrato(caminho_json):
    caminho_json = Path(caminho_json)

    with caminho_json.open('r', encoding='utf-8') as arquivo:
        dados = json.load(arquivo)

    caminho_modelo = Path(dados['modelo'])
    caminho_saida = Path(dados['saida'])

    if not caminho_modelo.is_file():
        raise FileNotFoundError(f'Modelo de contrato não encontrado: {caminho_modelo}')

    referencias = {
        str(codigo): '' if valor is None else str(valor)
        for codigo, valor in dados.get('referencias', {}).items()
    }

    referencias['{{VALOR_ALUGUEL_EXTENSO}}'] = valor_por_extenso(
        dados['valor_aluguel']
    )

    remover_se_vazio = set(dados.get('remover_se_vazio', []))

    documento = Document(str(caminho_modelo))

    processar_container(documento, referencias, remover_se_vazio)

    for secao in documento.sections:
        containers = [
            secao.header,
            secao.first_page_header,
            secao.even_page_header,
            secao.footer,
            secao.first_page_footer,
            secao.even_page_footer,
        ]

        for container in containers:
            processar_container(container, referencias, remover_se_vazio)

    texto_final = coletar_texto_documento(documento)
    placeholders_restantes = sorted(set(re.findall(r'\{\{[^{}]+\}\}', texto_final)))

    if placeholders_restantes:
        raise ValueError(
            'Existem campos do modelo que não foram preenchidos: ' +
            ', '.join(placeholders_restantes)
        )

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(caminho_saida))

    print(json.dumps({
        'sucesso': True,
        'arquivo': str(caminho_saida)
    }, ensure_ascii=False))


if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(
            'Uso: python gerar_contrato.py caminho/dados.json',
            file=sys.stderr
        )
        sys.exit(1)

    try:
        gerar_contrato(sys.argv[1])
    except Exception as erro:
        print(str(erro), file=sys.stderr)
        sys.exit(1)